"""
Document Text Extraction & Parsing Module.

Supports multi-format text extraction across:
- PDF (.pdf) via pdfplumber
- DOCX (.docx) via python-docx (paragraphs and tables)
- Plain text (.txt) with utf-8 / latin-1 fallback
- XML (.xml) via xml.etree.ElementTree with regex tag-stripping fallback

Features:
- Unified dispatcher: extract_text_any(source)
- Deduplicating resume loader: extract_resumes(folder_path)
  Prioritizes formats: .pdf > .docx > .txt > .xml
"""

from __future__ import annotations
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Union
from collections import defaultdict
import xml.etree.ElementTree as ET

import pdfplumber
from docx import Document


# Supported file formats and their priority for deduplication
SUPPORTED_EXTENSIONS: Set[str] = {".pdf", ".docx", ".txt", ".xml"}

FORMAT_PRIORITY: Dict[str, int] = {
    ".pdf": 0,
    ".docx": 1,
    ".txt": 2,
    ".xml": 3,
}


def extract_text_from_pdf(source: Union[str, Path, Any]) -> str:
    """
    Extracts all text from a given PDF file or buffer using pdfplumber.
    Supports file paths (str/Path), file-like objects (BytesIO), and Streamlit UploadedFiles.

    Args:
        source: Path to the PDF file or file-like object.

    Returns:
        Extracted raw text content as a string.
    """
    if hasattr(source, "getvalue"):
        file_obj = io.BytesIO(source.getvalue())
    elif hasattr(source, "read") and not isinstance(source, (str, Path)):
        file_obj = source
    elif isinstance(source, bytes):
        file_obj = io.BytesIO(source)
    else:
        pdf_path = Path(source)
        if not pdf_path.is_file():
            raise FileNotFoundError(f"PDF file not found: {pdf_path}")
        file_obj = str(pdf_path)

    extracted_pages = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text.strip())

    return "\n\n".join(extracted_pages).strip()


def extract_text_from_docx(source: Union[str, Path, Any]) -> str:
    """
    Extracts all text from a given DOCX file or buffer using python-docx.
    Loops through doc.paragraphs and doc.tables, joining all text.
    Supports file paths, file-like objects (BytesIO), and Streamlit UploadedFiles.

    Args:
        source: Path to the DOCX file or file-like object.

    Returns:
        Extracted text content as a string.
    """
    if hasattr(source, "getvalue"):
        file_obj = io.BytesIO(source.getvalue())
    elif hasattr(source, "read") and not isinstance(source, (str, Path)):
        file_obj = source
    elif isinstance(source, bytes):
        file_obj = io.BytesIO(source)
    else:
        docx_path = Path(source)
        if not docx_path.is_file():
            raise FileNotFoundError(f"DOCX file not found: {docx_path}")
        file_obj = str(docx_path)

    doc = Document(file_obj)
    text_parts: List[str] = []

    # 1. Extract paragraphs
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            text_parts.append(p_text)

    # 2. Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_cells_text = []
            for cell in row.cells:
                c_text = cell.text.strip()
                if c_text and (not row_cells_text or c_text != row_cells_text[-1]):
                    row_cells_text.append(c_text)
            if row_cells_text:
                text_parts.append("    ".join(row_cells_text))

    return "\n".join(text_parts).strip()


def extract_text_from_txt(source: Union[str, Path, Any]) -> str:
    """
    Extracts text from a plain text file or buffer (utf-8, fallback to latin-1 if decode fails).
    Supports file paths, file-like objects, and Streamlit UploadedFiles.

    Args:
        source: Path to the text file or file-like object.

    Returns:
        Extracted plain text content.
    """
    if hasattr(source, "getvalue"):
        raw_bytes = source.getvalue()
    elif hasattr(source, "read") and not isinstance(source, (str, Path)):
        content = source.read()
        if isinstance(content, str):
            return content.strip()
        raw_bytes = content
    elif isinstance(source, bytes):
        raw_bytes = source
    else:
        txt_path = Path(source)
        if not txt_path.is_file():
            raise FileNotFoundError(f"Text file not found: {txt_path}")
        try:
            return txt_path.read_text(encoding="utf-8").strip()
        except (UnicodeDecodeError, LookupError):
            return txt_path.read_text(encoding="latin-1", errors="ignore").strip()

    try:
        return raw_bytes.decode("utf-8").strip()
    except (UnicodeDecodeError, LookupError):
        return raw_bytes.decode("latin-1", errors="ignore").strip()


class ResumeDict(dict):
    """
    Dictionary mapping candidate resume names to extracted text.
    Allows lookups by full filename (e.g. 'resume.pdf') or by stem ('resume').
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._stem_map: Dict[str, str] = {}

    def __setitem__(self, key: str, value: str):
        super().__setitem__(key, value)
        stem = Path(str(key)).stem
        self._stem_map[stem] = key

    def __getitem__(self, key: str) -> str:
        if key in self:
            return super().__getitem__(key)
        stem = Path(str(key)).stem
        if stem in self._stem_map:
            return super().__getitem__(self._stem_map[stem])
        raise KeyError(key)

    def __contains__(self, key: object) -> bool:
        if super().__contains__(key):
            return True
        stem = Path(str(key)).stem
        return stem in self._stem_map


def _format_xml_tag_name(tag: str) -> str:
    """Converts XML tag names (e.g. 'technicalSkills', 'work_experience') to human-readable labels."""
    if "}" in tag:
        tag = tag.split("}", 1)[1]
    # camelCase to space-separated words
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", tag)
    # Underscores and hyphens to space
    s = re.sub(r"[_\-]+", " ", s).strip()
    return s.title()


def _walk_xml_element(elem: ET.Element, level: int = 0) -> List[str]:
    """
    Recursively walks an XML ElementTree node and extracts all visible text,
    tag-based structural headings, and key attributes down through all nested levels.
    """
    lines: List[str] = []
    tag_clean = elem.tag.split("}")[-1].lower() if "}" in elem.tag else elem.tag.lower()
    tag_title = _format_xml_tag_name(elem.tag)
    children = list(elem)

    # 1. Structural section header for container tags (except root document)
    if level > 0 and children:
        if tag_clean not in ("bullet", "bullets", "project", "certification", "applicant"):
            lines.append(f"\n{tag_title.upper()}:")
        elif tag_clean == "applicant":
            lines.append(f"\nCONTACT INFORMATION:")
    elif level > 0 and not children and tag_clean in (
        "summary", "education", "skills", "experience", "projects", "certifications"
    ):
        lines.append(f"\n{tag_title.upper()}:")

    # 2. Extract attributes (e.g. category="Languages", name="...", etc.)
    if elem.attrib:
        for attr_k, attr_v in elem.attrib.items():
            if attr_v and attr_v.strip():
                lines.append(f"{attr_v.strip()}:")

    # 3. Direct element text
    if elem.text and elem.text.strip():
        txt = elem.text.strip()
        if tag_clean == "bullet":
            lines.append(f"- {txt}")
        else:
            lines.append(txt)

    # 4. Recursively process all child elements
    for child in children:
        lines.extend(_walk_xml_element(child, level + 1))
        if child.tail and child.tail.strip():
            lines.append(child.tail.strip())

    return lines


def extract_text_from_xml(source: Union[str, Path, Any]) -> str:
    """
    Extracts visible text content from an XML file or buffer.
    Recursively traverses all XML elements, capturing structural headings (Education,
    Technical Skills, Experience, Projects, Certifications, Contact), element attributes,
    and text content across all hierarchy levels.
    Falls back to regex tag-stripping if XML is malformed or unstructured.

    Args:
        source: Path to the XML file or file-like object.

    Returns:
        Extracted visible text content.
    """
    raw_xml = extract_text_from_txt(source)
    if not raw_xml:
        return ""

    try:
        root = ET.fromstring(raw_xml.encode("utf-8"))
        extracted_lines = _walk_xml_element(root)
        full_text = "\n".join(extracted_lines).strip()
        if full_text:
            return full_text
    except Exception:
        pass

    # Fallback: strip XML/HTML tags with regex
    cleaned = re.sub(r"<[^>]+>", " ", raw_xml)
    lines = [re.sub(r"\s+", " ", line).strip() for line in cleaned.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def extract_text_any(source: Union[str, Path, Any]) -> str:
    """
    Dispatcher function that checks the file extension (.pdf, .docx, .txt, .xml)
    and calls the appropriate extractor, raising a clear ValueError for unsupported types.

    Args:
        source: File path, file-like object, or Streamlit UploadedFile.

    Returns:
        Extracted text content.
    """
    if hasattr(source, "name"):
        ext = Path(source.name).suffix.lower()
    elif isinstance(source, (str, Path)):
        ext = Path(source).suffix.lower()
    else:
        raise ValueError(f"Cannot determine file extension for source of type: {type(source).__name__}")

    if ext == ".pdf":
        return extract_text_from_pdf(source)
    elif ext == ".docx":
        return extract_text_from_docx(source)
    elif ext in (".txt", ".text"):
        return extract_text_from_txt(source)
    elif ext == ".xml":
        return extract_text_from_xml(source)
    else:
        supported_str = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported formats are: {supported_str}"
        )


def extract_jd(source: Union[str, Path, Any]) -> str:
    """
    Extracts raw Job Description text from a file (.pdf, .docx, .txt, .xml, .md) or uploaded buffer.

    Args:
        source: Path to the Job Description file or UploadedFile.

    Returns:
        Extracted raw JD text.
    """
    try:
        return extract_text_any(source)
    except ValueError:
        # Fallback for markdown (.md) or unknown text formats
        return extract_text_from_txt(source)


def get_dedup_mapping(folder_path: Union[str, Path]) -> Dict[str, Dict[str, Any]]:
    """
    Computes the deduplication mapping for all candidate resume files in folder_path.
    Groups files by base filename (stem) and selects one file using priority:
    .pdf > .docx > .txt > .xml.

    Args:
        folder_path: Directory path containing resume files.

    Returns:
        Dict mapping candidate stem to {'kept': str, 'dropped': List[str]}.
    """
    folder = Path(folder_path)
    if not folder.is_dir():
        raise NotADirectoryError(f"Resume folder not found: {folder}")

    all_files = [
        f for f in folder.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if not all_files:
        all_files = [
            f for f in folder.rglob("*")
            if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
        ]

    candidates_files: Dict[str, List[Path]] = defaultdict(list)
    for file_path in all_files:
        candidates_files[file_path.stem].append(file_path)

    mapping: Dict[str, Dict[str, Any]] = {}
    for stem in sorted(candidates_files.keys()):
        files_list = candidates_files[stem]
        chosen_file = min(
            files_list,
            key=lambda f: FORMAT_PRIORITY.get(f.suffix.lower(), 99),
        )
        dropped_files = [f for f in files_list if f != chosen_file]
        mapping[stem] = {
            "kept": chosen_file.name,
            "dropped": [f.name for f in dropped_files],
        }

    return mapping


def extract_resumes(folder_path: Union[str, Path], verbose: bool = True) -> Dict[str, str]:
    """
    Extracts text from all supported resume files (.pdf, .docx, .txt, .xml) in folder_path.
    If the same candidate appears as multiple files with different formats (e.g. Resume.pdf,
    Resume.docx, Resume.txt, Resume.xml), groups files by their base filename (stripping extension)
    and keeps only one file per candidate, preferring .pdf > .docx > .txt > .xml in that priority order.

    Prints the complete deduplication mapping showing which files were kept and which were dropped.

    Args:
        folder_path: Directory path containing resume files.
        verbose: Whether to print the deduplication report (default: True).

    Returns:
        ResumeDict mapping chosen filename (e.g. 'Resume.pdf') to extracted text.
    """
    folder = Path(folder_path)
    if not folder.is_dir():
        raise NotADirectoryError(f"Resume folder not found: {folder}")

    all_files = [
        f for f in folder.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if not all_files:
        all_files = [
            f for f in folder.rglob("*")
            if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
        ]

    # Group files by base filename (stem)
    candidates_files: Dict[str, List[Path]] = defaultdict(list)
    for file_path in all_files:
        candidates_files[file_path.stem].append(file_path)

    resumes: Dict[str, str] = ResumeDict()
    dedup_report: List[tuple] = []

    # Sort candidates by name and choose highest priority file per candidate
    for stem in sorted(candidates_files.keys()):
        files_list = candidates_files[stem]
        chosen_file = min(
            files_list,
            key=lambda f: FORMAT_PRIORITY.get(f.suffix.lower(), 99),
        )
        dropped_files = [f for f in files_list if f != chosen_file]
        if dropped_files:
            dedup_report.append((stem, chosen_file, dropped_files))

        try:
            resumes[chosen_file.name] = extract_text_any(chosen_file)
        except Exception as e:
            print(f"Warning: Failed to parse '{chosen_file.name}': {e}")
            resumes[chosen_file.name] = ""

    # Print the deduplication mapping
    if verbose:
        print("=" * 80)
        print("RESUME DEDUPLICATION REPORT")
        print("=" * 80)
        print(f"Directory:              {folder}")
        print(f"Total files scanned:    {len(all_files)}")
        print(f"Unique candidate stems: {len(candidates_files)}")
        print(f"Priority order:         .pdf > .docx > .txt > .xml\n")
        if dedup_report:
            print(f"Duplicates Detected & Resolved ({len(dedup_report)} candidate groups):")
            print("-" * 80)
            for stem, chosen_file, dropped_files in dedup_report:
                dropped_str = ", ".join(f.name for f in dropped_files)
                print(f"Candidate: '{stem}'")
                print(f"  [KEPT]    {chosen_file.name} ({chosen_file.suffix.upper()})")
                print(f"  [DROPPED] {dropped_str}\n")
            print("-" * 80)
            total_dropped = sum(len(d) for _, _, d in dedup_report)
            print(
                f"Deduplication summary: {total_dropped} redundant files dropped across "
                f"{len(dedup_report)} candidate groups."
            )
        else:
            print("No duplicate formats detected among files.")
        print(f"Final active candidates retained: {len(resumes)}")
        print("=" * 80 + "\n")

    return resumes


def main():
    """Runs verification of multi-format reading and deduplication."""
    base_dir = Path(__file__).resolve().parent
    new_folder = base_dir / "test_data" / "dummy_resumes" / "Dummy Resumes"
    if not new_folder.exists():
        new_folder = base_dir / "test_data" / "dummy_resumes"
    if not new_folder.exists():
        new_folder = Path(r"C:\Users\SES\Downloads\Dummy Resumes")
    if not new_folder.exists():
        new_folder = base_dir / "test_data" / "resumes"

    print("=" * 80)
    print("PARSER MODULE - MULTI-FORMAT EXTRACTION & DEDUPLICATION TEST")
    print("=" * 80)
    print(f"Target Directory: {new_folder}")
    print(f"Priority Order:   .pdf > .docx > .txt > .xml\n")

    resumes = extract_resumes(new_folder)
    print(f"Total Unique Candidates Extracted: {len(resumes)}")
    print("-" * 80)
    print(f"{'Candidate / Filename':<55} | {'Length':<10} | Format")
    print("-" * 80)
    for filename, text in resumes.items():
        ext = Path(filename).suffix.upper()
        print(f"{filename:<55} | {len(text):<6} chars | {ext}")

    print("-" * 80)
    print(f"Extraction and deduplication completed successfully ({len(resumes)} candidates).")
    print("=" * 80)


if __name__ == "__main__":
    main()
